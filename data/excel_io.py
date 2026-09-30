"""The coordinator's workbook in and out.

`export_template(u)` writes a campus as the .xlsx she already keeps: courses,
sections, rooms, buildings, instructors and (optionally) student groups.
`load_workbook(data)` reads one back into a `University` and a validation
report. No LLM touches this file: every cell is parsed by code, and the same
file always produces the same campus.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook as _open_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.models import (
    DAYS,
    DAYS_AR,
    FIRST_HOUR,
    N_DAYS,
    N_PERIODS,
    Building,
    Course,
    Instructor,
    Meeting,
    Room,
    Section,
    Student,
    University,
    slot_day,
    slot_id,
    slot_period,
)

DEFAULT_ENROLLMENT = 20
DEFAULT_WALK = 5
MEETINGS_PER_SECTION = 2

# sheet key -> (Arabic title, required)
SHEETS = {
    "Courses": ("المقررات", True),
    "Sections": ("الشعب", True),
    "StudentGroups": ("مجموعات الطلاب", False),
    "Rooms": ("القاعات", True),
    "Buildings": ("المباني", False),
    "WalkMinutes": ("دقائق المشي", False),
    "Instructors": ("المدرسون", False),
}

# column key -> Arabic header. Headers are written "عربي (key)"; either half
# is accepted on the way back in, so a coordinator can retype them in Arabic.
COLUMNS = {
    "Courses": [
        ("course_code", "رمز المقرر"),
        ("name_ar", "اسم المقرر"),
        ("name_en", "الاسم بالإنجليزية"),
        ("department", "القسم"),
        ("level", "المستوى"),
        ("needs_lab", "يحتاج معمل"),
    ],
    "Sections": [
        ("section_id", "رقم الشعبة"),
        ("course_code", "رمز المقرر"),
        ("instructor_name", "المدرس"),
        ("enrollment", "عدد الطلاب"),
        ("cohorts", "الدفعات"),
    ],
    "StudentGroups": [
        ("group_id", "المجموعة"),
        ("department", "القسم"),
        ("level", "المستوى"),
        ("student_count", "عدد الطلاب"),
        ("section_ids", "الشعب"),
        ("repeater", "معيد"),
        ("needs_accessibility", "يحتاجون قاعة مهيأة"),
    ],
    "Rooms": [
        ("room_id", "القاعة"),
        ("building", "المبنى"),
        ("capacity", "السعة"),
        ("accessible", "مهيأة"),
        ("kind", "النوع"),
    ],
    "Buildings": [
        ("building_id", "رمز المبنى"),
        ("name_ar", "اسم المبنى"),
        ("name_en", "الاسم بالإنجليزية"),
    ],
    "WalkMinutes": [
        ("from_building", "من مبنى"),
        ("to_building", "إلى مبنى"),
        ("minutes", "الدقائق"),
    ],
    "Instructors": [
        ("instructor_name", "المدرس"),
        ("name_en", "الاسم بالإنجليزية"),
        ("unavailable", "أوقات غير متاح"),
    ],
}

DAY_TOKENS = {
    **{d.lower(): i for i, d in enumerate(DAYS)},
    **{d[:3].lower(): i for i, d in enumerate(DAYS)},
    **{d: i for i, d in enumerate(DAYS_AR)},
    "اثنين": 1, "الإثنين": 1, "أحد": 0, "ثلاثاء": 2, "أربعاء": 3, "خميس": 4,
    "الاربعاء": 3, "احد": 0, "اربعاء": 3,
}
DAY_SHORT = [d[:3] for d in DAYS]

YES = {"yes", "y", "true", "1", "نعم", "اي", "أي", "ايوه", "صح", "x", "✓"}
NO = {"no", "n", "false", "0", "لا", "", "none"}
LAB_WORDS = {"lab", "laboratory", "معمل", "مختبر"}

ARABIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


# --- report ---------------------------------------------------------------


@dataclass
class Report:
    errors: list[dict] = field(default_factory=list)
    warnings: list[dict] = field(default_factory=list)
    notes: list[dict] = field(default_factory=list)
    summary: dict = field(default_factory=dict)

    def error(self, sheet: str | None, row: int | None, en: str, ar: str) -> None:
        self.errors.append({"sheet": sheet, "row": row, "en": en, "ar": ar})

    def warn(self, sheet: str | None, row: int | None, en: str, ar: str) -> None:
        self.warnings.append({"sheet": sheet, "row": row, "en": en, "ar": ar})

    def note(self, en: str, ar: str) -> None:
        self.notes.append({"en": en, "ar": ar})

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_dict(self) -> dict:
        return {
            "ok": self.ok,
            "errors": self.errors,
            "warnings": self.warnings,
            "notes": self.notes,
            "summary": self.summary,
        }


class WorkbookError(ValueError):
    """The file cannot be read at all. Carries the report so far."""

    def __init__(self, report: Report) -> None:
        super().__init__(report.errors[0]["en"] if report.errors else "invalid workbook")
        self.report = report


# --- cell parsing -----------------------------------------------------------


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).translate(ARABIC_DIGITS).strip()


def _int(v) -> int | None:
    t = _text(v)
    if not t:
        return None
    try:
        f = float(t)
    except ValueError:
        return None
    return int(f) if f.is_integer() else None


def _yes(v, default: bool = False) -> bool | None:
    t = _text(v).lower()
    if not t:
        return default
    if t in YES:
        return True
    if t in NO:
        return False
    return None


def _list(v) -> list[str]:
    return [x.strip() for x in re.split(r"[,،;\s]+", _text(v)) if x.strip()]


def parse_unavailable(text: str) -> set[int]:
    """"Tue 14-16; Sun all" -> slot ids. Ranges are end-exclusive hours:
    14-16 blocks the 14:00 and 15:00 slots. Raises ValueError on nonsense."""
    out: set[int] = set()
    for part in re.split(r"[;؛\n]+", _text(text)):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\S+)\s*(.*)$", part)
        day_word, rest = m.group(1), m.group(2).strip()
        day = DAY_TOKENS.get(day_word.lower().rstrip(":"))
        if day is None:
            raise ValueError(f"unknown day '{day_word}'")
        if not rest or rest.lower() in {"all", "كامل", "طول اليوم", "كل اليوم"}:
            out |= {slot_id(day, p) for p in range(N_PERIODS)}
            continue
        for rng in re.split(r"[,،]+", rest):
            rng = rng.strip()
            if not rng:
                continue
            hm = re.match(r"^(\d{1,2})(?::00)?\s*(?:-|–|to|إلى|الى)?\s*(\d{1,2})?(?::00)?$", rng)
            if not hm:
                raise ValueError(f"cannot read the hours '{rng}'")
            start = int(hm.group(1))
            end = int(hm.group(2)) if hm.group(2) else start + 1
            if not (FIRST_HOUR <= start < end <= FIRST_HOUR + N_PERIODS):
                raise ValueError(
                    f"hours {start}-{end} are outside "
                    f"{FIRST_HOUR}-{FIRST_HOUR + N_PERIODS}"
                )
            out |= {slot_id(day, h - FIRST_HOUR) for h in range(start, end)}
    return out


def format_unavailable(slots: set[int]) -> str:
    """The inverse of parse_unavailable: "Sun 08-09, 13-15; Tue all"."""
    parts = []
    for d in range(N_DAYS):
        ps = sorted(slot_period(s) for s in slots if slot_day(s) == d)
        if not ps:
            continue
        if ps == list(range(N_PERIODS)):
            parts.append(f"{DAY_SHORT[d]} all")
            continue
        runs: list[list[int]] = []
        for p in ps:
            if runs and p == runs[-1][1]:
                runs[-1][1] = p + 1
            else:
                runs.append([p, p + 1])
        parts.append(
            DAY_SHORT[d] + " " + ", ".join(
                f"{FIRST_HOUR + a:02d}-{FIRST_HOUR + b:02d}" for a, b in runs
            )
        )
    return "; ".join(parts)


# --- reading ----------------------------------------------------------------


def _norm_header(h) -> str:
    return re.sub(r"\s+", " ", _text(h)).strip().lower()


def _find_sheet(wb, key: str):
    ar = SHEETS[key][0]
    for ws in wb.worksheets:
        title = ws.title.strip()
        if key.lower() in title.lower().replace(" ", "") or ar in title:
            return ws
    return None


def _rows(ws, key: str, report: Report) -> list[tuple[int, dict]]:
    """[(excel row number, {column key: value})], skipping blank rows. Extra
    columns are kept under their header text (the walking matrix uses this)."""
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        return []
    known = {}
    for k, ar in COLUMNS[key]:
        known[k.lower()] = k
        known[ar] = k
    keys: list[str | None] = []
    for h in header:
        t = _norm_header(h)
        m = re.search(r"\(([a-z_]+)\)\s*$", t)
        if m and m.group(1) in known:
            keys.append(known[m.group(1)])
        elif t in known:
            keys.append(known[t])
        elif _text(h):
            keys.append(_text(h))
        else:
            keys.append(None)
    required = {k for k, _ in COLUMNS[key]} - OPTIONAL_COLUMNS[key]
    missing = [k for k in required if k not in keys]
    for k in missing:
        ar = dict(COLUMNS[key])[k]
        report.error(
            key, 1,
            f"column '{k}' is missing",
            f"العمود «{ar}» غير موجود",
        )
    out = []
    for i, row in enumerate(it, start=2):
        if row is None or all(_text(v) == "" for v in row):
            continue
        out.append((i, {k: v for k, v in zip(keys, row) if k is not None}))
    return out


OPTIONAL_COLUMNS = {
    "Courses": {"name_en", "needs_lab"},
    "Sections": {"enrollment", "cohorts"},
    "StudentGroups": {"group_id", "repeater", "needs_accessibility"},
    "Rooms": {"accessible", "kind"},
    "Buildings": {"name_ar", "name_en"},
    "WalkMinutes": set(),
    "Instructors": {"name_en", "unavailable"},
}


def load_workbook(data: bytes) -> tuple[University | None, dict]:
    """-> (University or None when there are errors, report dict)."""
    report = Report()
    try:
        wb = _open_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:  # not an xlsx at all
        report.error(None, None, f"not a readable .xlsx file ({type(e).__name__})",
                     "الملف ليس ملف Excel صالحًا (.xlsx)")
        return None, report.as_dict()

    sheets = {k: _find_sheet(wb, k) for k in SHEETS}
    for k, (ar, required) in SHEETS.items():
        if required and sheets[k] is None:
            report.error(k, None, f"sheet '{k}' is missing", f"ورقة «{ar}» غير موجودة")
    if not report.ok:
        return None, report.as_dict()

    rows = {k: (_rows(ws, k, report) if ws is not None else None) for k, ws in sheets.items()}
    if not report.ok:
        return None, report.as_dict()

    u = _build(rows, report)
    return (u if report.ok else None), report.as_dict()


def _build(rows: dict, report: Report) -> University | None:
    # --- buildings and rooms
    buildings: list[Building] = []
    seen_b: set[str] = set()
    for r, row in rows["Buildings"] or []:
        bid = _text(row.get("building_id"))
        if not bid:
            report.error("Buildings", r, "building id is empty", "رمز المبنى فارغ")
            continue
        if bid in seen_b:
            report.error("Buildings", r, f"building {bid} is listed twice",
                         f"المبنى {bid} مكرر")
            continue
        seen_b.add(bid)
        name_ar = _text(row.get("name_ar")) or bid
        buildings.append(Building(bid, _text(row.get("name_en")) or name_ar, name_ar))

    rooms: list[Room] = []
    seen_r: set[str] = set()
    for r, row in rows["Rooms"]:
        rid = _text(row.get("room_id"))
        bid = _text(row.get("building"))
        cap = _int(row.get("capacity"))
        acc = _yes(row.get("accessible"))
        kind_raw = _text(row.get("kind")).lower()
        if not rid:
            report.error("Rooms", r, "room id is empty", "رمز القاعة فارغ")
            continue
        if rid in seen_r:
            report.error("Rooms", r, f"room {rid} is listed twice", f"القاعة {rid} مكررة")
            continue
        seen_r.add(rid)
        if cap is None or cap <= 0:
            report.error("Rooms", r, f"room {rid}: capacity must be a positive number",
                         f"القاعة {rid}: السعة يجب أن تكون رقمًا موجبًا")
            continue
        if acc is None:
            report.error("Rooms", r, f"room {rid}: 'accessible' must be yes or no",
                         f"القاعة {rid}: «مهيأة» يجب أن تكون نعم أو لا")
            continue
        if not bid:
            report.error("Rooms", r, f"room {rid} has no building",
                         f"القاعة {rid} بدون مبنى")
            continue
        if bid not in seen_b:
            if rows["Buildings"] is not None:
                report.warn("Rooms", r, f"building {bid} is not in the Buildings sheet; added",
                            f"المبنى {bid} غير موجود في ورقة المباني؛ أُضيف")
            seen_b.add(bid)
            buildings.append(Building(bid, bid, bid))
        rooms.append(Room(rid, bid, cap, acc, "lab" if kind_raw in LAB_WORDS else "lecture"))

    walk = _walk(rows, buildings, report)

    # --- instructors: sheet order first, then first appearance in Sections
    instructors: list[Instructor] = []
    inst_by_name: dict[str, Instructor] = {}

    def add_instructor(name: str, name_en: str = "") -> Instructor:
        inst = Instructor(f"I{len(instructors) + 1:02d}", name, name_en)
        instructors.append(inst)
        inst_by_name[name] = inst
        return inst

    for r, row in rows["Instructors"] or []:
        name = _text(row.get("instructor_name"))
        if not name:
            report.error("Instructors", r, "instructor name is empty", "اسم المدرس فارغ")
            continue
        if name in inst_by_name:
            report.error("Instructors", r, f"{name} is listed twice", f"{name} مكرر")
            continue
        inst = add_instructor(name, _text(row.get("name_en")))
        try:
            inst.unavailable_slots = parse_unavailable(row.get("unavailable"))
        except ValueError as e:
            report.error("Instructors", r, f"{name}: {e} (write e.g. 'Tue 14-16; Sun all')",
                         f"{name}: تعذّرت قراءة الأوقات (اكتبي مثلًا: Tue 14-16; Sun all)")

    # --- courses
    courses: list[Course] = []
    course_by_code: dict[str, Course] = {}
    for r, row in rows["Courses"]:
        code = _text(row.get("course_code"))
        name_ar = _text(row.get("name_ar"))
        dept = _text(row.get("department"))
        level = _int(row.get("level"))
        lab = _yes(row.get("needs_lab"))
        if not code:
            report.error("Courses", r, "course code is empty", "رمز المقرر فارغ")
            continue
        if code in course_by_code:
            report.error("Courses", r, f"course {code} is listed twice", f"المقرر {code} مكرر")
            continue
        if not dept:
            report.error("Courses", r, f"course {code} has no department",
                         f"المقرر {code} بدون قسم")
            continue
        if level is None or level < 1:
            report.error("Courses", r, f"course {code}: level must be a whole number ≥ 1",
                         f"المقرر {code}: المستوى يجب أن يكون رقمًا صحيحًا ≥ ١")
            continue
        if lab is None:
            report.error("Courses", r, f"course {code}: 'needs_lab' must be yes or no",
                         f"المقرر {code}: «يحتاج معمل» يجب أن تكون نعم أو لا")
            continue
        c = Course(code, _text(row.get("name_en")) or name_ar or code,
                   name_ar or _text(row.get("name_en")) or code, dept, level, lab)
        courses.append(c)
        course_by_code[code] = c

    # --- sections
    sections: list[Section] = []
    section_rows: dict[str, int] = {}
    cohorts_given: dict[str, list[str]] = {}
    for r, row in rows["Sections"]:
        sid = _text(row.get("section_id"))
        code = _text(row.get("course_code"))
        iname = _text(row.get("instructor_name"))
        enroll = _int(row.get("enrollment"))
        if not sid:
            sid = f"S{len(sections) + 1:02d}"
        if sid in section_rows:
            report.error("Sections", r, f"section {sid} is listed twice", f"الشعبة {sid} مكررة")
            continue
        if code not in course_by_code:
            report.error("Sections", r, f"section {sid}: unknown course code '{code}'",
                         f"الشعبة {sid}: رمز المقرر «{code}» غير موجود في ورقة المقررات")
            continue
        if not iname:
            report.error("Sections", r, f"section {sid} has no instructor",
                         f"الشعبة {sid} بدون مدرس")
            continue
        if _text(row.get("enrollment")) == "":
            enroll = DEFAULT_ENROLLMENT
        if enroll is None or enroll <= 0:
            report.error("Sections", r, f"section {sid}: enrollment must be a positive number",
                         f"الشعبة {sid}: عدد الطلاب يجب أن يكون رقمًا موجبًا")
            continue
        inst = inst_by_name.get(iname)
        if inst is None:
            if rows["Instructors"] is not None:
                report.warn("Sections", r, f"{iname} is not in the Instructors sheet; added "
                            "as always available",
                            f"{iname} غير موجود في ورقة المدرسين؛ أُضيف متاحًا طوال الأسبوع")
            inst = add_instructor(iname)
        course = course_by_code[code]
        sections.append(Section(sid, code, inst.id, enroll, course.department, course.level))
        section_rows[sid] = r
        if _list(row.get("cohorts")):
            cohorts_given[sid] = _list(row.get("cohorts"))

    if not sections:
        report.error("Sections", None, "no sections to schedule", "لا توجد شعب للجدولة")
        return None

    # --- student groups: given, or derived the way the coordinator works
    groups = _groups(rows, courses, sections, section_rows, report)

    students: list[Student] = []
    serving: dict[str, set[str]] = {s.id: set() for s in sections}
    for g in groups:
        for k in range(g["count"]):
            n = len(students) + 1
            students.append(Student(
                id=f"ST{n:03d}",
                name=f"{g['id']} · {k + 1}",
                department=g["department"],
                level=g["level"],
                section_ids=list(g["sections"]),
                needs_accessibility=k < g["access"],
                is_repeater=g["repeater"],
                name_en=f"{g['id']} · {k + 1}",
            ))
        if not g["repeater"]:
            for sid in g["sections"]:
                serving[sid].add(f"{g['department']}-L{g['level']}")
    for s in sections:
        if s.id in cohorts_given:
            serving[s.id] = set(cohorts_given[s.id])
        elif not serving[s.id]:
            serving[s.id] = {s.cohort}

    meetings = [
        Meeting(f"{s.id}-m{k + 1}", s.id, k)
        for s in sections
        for k in range(MEETINGS_PER_SECTION)
    ]
    u = University(
        buildings=buildings,
        rooms=rooms,
        instructors=instructors,
        courses=courses,
        sections=sections,
        meetings=meetings,
        students=students,
        walk=walk,
        serving_cohorts=serving,
    )
    _check_feasible(u, section_rows, report)
    report.summary = {
        "levels": len({(c.department, c.level) for c in courses}),
        "departments": len({c.department for c in courses}),
        "courses": len(courses),
        "sections": len(sections),
        "meetings": len(meetings),
        "students": len(students),
        "student_groups": len(groups),
        "repeaters": sum(1 for s in students if s.is_repeater),
        "needs_accessibility": sum(1 for s in students if s.needs_accessibility),
        "instructors": len(instructors),
        "rooms": len(rooms),
        "buildings": len(buildings),
    }
    return u


def _walk(rows: dict, buildings: list[Building], report: Report) -> dict:
    ids = [b.id for b in buildings]
    given: dict[tuple[str, str], int] = {}
    # a matrix: extra columns on the Buildings sheet named after building ids
    for r, row in rows["Buildings"] or []:
        a = _text(row.get("building_id"))
        for b in ids:
            v = _int(row.get(b))
            if v is not None and a != b:
                given[(a, b)] = v
    # or a long table
    for r, row in rows["WalkMinutes"] or []:
        a, b = _text(row.get("from_building")), _text(row.get("to_building"))
        v = _int(row.get("minutes"))
        if a not in ids or b not in ids or v is None or v < 0:
            report.warn("WalkMinutes", r, "row ignored: unknown building or bad minutes",
                        "تم تجاهل الصف: مبنى غير معروف أو دقائق غير صحيحة")
            continue
        given[(a, b)] = v
    walk: dict[tuple[str, str], int] = {}
    defaulted = False
    for a in ids:
        for b in ids:
            if a == b:
                walk[(a, b)] = 0
            elif (a, b) in given:
                walk[(a, b)] = given[(a, b)]
            elif (b, a) in given:
                walk[(a, b)] = given[(b, a)]
            else:
                walk[(a, b)] = DEFAULT_WALK
                defaulted = True
    if defaulted and len(ids) > 1:
        report.note(
            f"walking time not given for some buildings: {DEFAULT_WALK} minutes assumed",
            f"لم تُحدد دقائق المشي بين بعض المباني: افترضنا {DEFAULT_WALK} دقائق",
        )
    return walk


def _groups(rows, courses, sections, section_rows, report) -> list[dict]:
    known = {s.id for s in sections}
    if rows["StudentGroups"] is not None:
        out = []
        for r, row in rows["StudentGroups"]:
            gid = _text(row.get("group_id")) or f"G{len(out) + 1:02d}"
            dept = _text(row.get("department"))
            level = _int(row.get("level"))
            count = _int(row.get("student_count"))
            sids = _list(row.get("section_ids"))
            rep = _yes(row.get("repeater"))
            acc = _int(row.get("needs_accessibility")) if _text(
                row.get("needs_accessibility")) else 0
            bad = False
            if not dept or level is None or level < 1:
                report.error("StudentGroups", r, f"group {gid}: department and level are required",
                             f"المجموعة {gid}: القسم والمستوى مطلوبان")
                bad = True
            if count is None or count <= 0:
                report.error("StudentGroups", r, f"group {gid}: student count must be positive",
                             f"المجموعة {gid}: عدد الطلاب يجب أن يكون موجبًا")
                bad = True
            for sid in sids:
                if sid not in known:
                    report.error("StudentGroups", r,
                                 f"group {gid}: section {sid} is not in the Sections sheet",
                                 f"المجموعة {gid}: الشعبة {sid} غير موجودة في ورقة الشعب")
                    bad = True
            if not sids:
                report.error("StudentGroups", r, f"group {gid} takes no sections",
                             f"المجموعة {gid} بدون شعب")
                bad = True
            if rep is None or acc is None or acc < 0 or (count and acc > count):
                report.error("StudentGroups", r,
                             f"group {gid}: 'repeater' must be yes/no and accessibility a "
                             "count no larger than the group",
                             f"المجموعة {gid}: «معيد» نعم/لا، وعدد ذوي الإعاقة لا يتجاوز المجموعة")
                bad = True
            if not bad:
                out.append({"id": gid, "department": dept, "level": level, "count": count,
                            "sections": sids, "repeater": rep, "access": acc})
        taken = {sid for g in out for sid in g["sections"]}
        for s in sections:
            if s.id not in taken:
                report.warn("Sections", section_rows[s.id],
                            f"section {s.id} is not taken by any student group",
                            f"الشعبة {s.id} لا تأخذها أي مجموعة طلاب")
        return out

    # Derived: per (department, level), k groups where k is the most sections
    # any course in the level has; group g takes section g mod n of each course.
    by_course: dict[str, list[Section]] = {}
    for s in sections:
        by_course.setdefault(s.course_id, []).append(s)
    levels: dict[tuple[str, int], list[Course]] = {}
    for c in courses:
        if c.id in by_course:
            levels.setdefault((c.department, c.level), []).append(c)
    out = []
    for (dept, level), cs in levels.items():
        k = max(len(by_course[c.id]) for c in cs)
        total = max(sum(s.enrollment for s in by_course[c.id]) for c in cs)
        for g in range(k):
            out.append({
                "id": f"{dept}-L{level}-G{g + 1}",
                "department": dept,
                "level": level,
                "count": total // k + (1 if g < total % k else 0),
                "sections": [by_course[c.id][g % len(by_course[c.id])].id for c in cs],
                "repeater": False,
                "access": 0,
            })
    report.note(
        f"no StudentGroups sheet: built {len(out)} groups from the study plan — for each "
        "level, one group per section of its largest course; group g takes section "
        "g mod n of every course in the level, and the level's students are split "
        "evenly across its groups",
        f"لا توجد ورقة مجموعات الطلاب: كوّنّا {len(out)} مجموعة من الخطة الدراسية — لكل "
        "مستوى مجموعة لكل شعبة من أكبر مقرراته، وتأخذ المجموعات شعب كل مقرر بالتناوب، "
        "ويُوزّع طلاب المستوى بالتساوي",
    )
    return out


def _check_feasible(u: University, section_rows: dict, report: Report) -> None:
    """The cheap checks that would otherwise surface as a failed build."""
    need_access = u.sections_needing_accessible()
    labs = [r for r in u.rooms if r.kind == "lab"]
    for s in u.sections:
        row = section_rows.get(s.id)
        course = u.course_by_id[s.course_id]
        pool = labs if course.needs_lab else u.rooms
        fits = [r for r in pool if r.capacity >= s.enrollment]
        if not fits:
            what = "lab" if course.needs_lab else "room"
            what_ar = "معمل" if course.needs_lab else "قاعة"
            report.error("Sections", row,
                         f"section {s.id}: no {what} seats {s.enrollment} students",
                         f"الشعبة {s.id}: لا يوجد {what_ar} يتسع لـ {s.enrollment} طالبًا")
        elif s.id in need_access and not any(r.accessible for r in fits):
            report.error("Sections", row,
                         f"section {s.id}: a student needs an accessible room and none "
                         f"seats {s.enrollment}",
                         f"الشعبة {s.id}: فيها طالب يحتاج قاعة مهيأة ولا توجد قاعة مهيأة "
                         f"تتسع لـ {s.enrollment}")
        taking = len(u.students_of_section[s.id])
        if taking > s.enrollment:
            report.warn("Sections", row,
                        f"section {s.id}: {taking} students in groups but enrollment "
                        f"is {s.enrollment}; rooms are sized by enrollment",
                        f"الشعبة {s.id}: {taking} طالبًا في المجموعات والعدد المسجل "
                        f"{s.enrollment}؛ القاعات تُختار حسب العدد المسجل")
    slots = N_DAYS * N_PERIODS
    if len(u.meetings) > len(u.rooms) * slots:
        report.error(None, None,
                     f"{len(u.meetings)} meetings do not fit in {len(u.rooms)} rooms × "
                     f"{slots} hours",
                     f"{len(u.meetings)} محاضرة لا تتسع لها {len(u.rooms)} قاعة × {slots} ساعة")
    for inst in u.instructors:
        n = sum(len(u.meetings_of_section[s.id]) for s in u.sections
                if s.instructor_id == inst.id)
        free = slots - len(inst.unavailable_slots)
        if n > free:
            report.error("Instructors", None,
                         f"{inst.name} teaches {n} meetings but is available for {free} hours",
                         f"{inst.name} لديه {n} محاضرة ومتاح {free} ساعة فقط")


# --- writing ----------------------------------------------------------------

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
GUIDE = [
    ("جدول — قالب بيانات الجدول الدراسي", "Jadwal — timetable data template"),
    ("", ""),
    ("املئي الأوراق كما تحضّرينها اليوم. أسماء الأعمدة بالعربي أو بالإنجليزي.",
     "Fill the sheets as you prepare them today. Headers may be Arabic or English."),
    ("المقررات: رمز المقرر، الاسم، القسم، المستوى، يحتاج معمل (نعم/لا).",
     "Courses: code, name, department, level, needs_lab (yes/no)."),
    ("الشعب: رقم الشعبة، رمز المقرر، المدرس، عدد الطلاب (افتراضيًا 20)، الدفعات (اختياري).",
     "Sections: id, course code, instructor, enrollment (default 20), cohorts (optional)."),
    ("مجموعات الطلاب اختيارية: إن حذفتِها نكوّنها من الخطة وعدد الشعب.",
     "StudentGroups is optional: without it, groups are built from the plan and section counts."),
    ("القاعات: رمز القاعة، المبنى، السعة، مهيأة (نعم/لا)، النوع (lecture/lab).",
     "Rooms: id, building, capacity, accessible (yes/no), kind (lecture/lab)."),
    ("المباني: مع مصفوفة دقائق المشي؛ إن لم تُذكر نفترض 5 دقائق.",
     "Buildings: with a walking-minutes matrix; 5 minutes assumed if missing."),
    ("المدرسون: أوقات غير متاح مثل: Tue 14-16; Sun all",
     "Instructors: unavailable, e.g. Tue 14-16; Sun all"),
    ("كل شعبة لها محاضرتان في يومين مختلفين، ساعة لكل محاضرة، من 8 إلى 4.",
     "Every section meets twice, on different days, one hour each, 08:00–16:00."),
]


def _sheet(wb: Workbook, key: str, header_keys: list[str], rows: list[list]):
    ws = wb.create_sheet(f"{SHEETS[key][0]} ({key})")
    ws.sheet_view.rightToLeft = True
    labels = dict(COLUMNS.get(key, []))
    ws.append([f"{labels[k]} ({k})" if k in labels else k for k in header_keys])
    for cell in ws[1]:
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, k in enumerate(header_keys):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(
            14, len(labels.get(k, k)) + len(k) + 6
        )
    ws.freeze_panes = "A2"
    return ws


def _yn(v: bool) -> str:
    return "yes" if v else "no"


def export_template(u: University) -> bytes:
    """The campus as the coordinator's workbook. Loading it back gives the same
    campus: same sections, enrolments, rooms, availability and students."""
    wb = Workbook()
    guide = wb.active
    guide.title = "اقرأني (Guide)"
    guide.sheet_view.rightToLeft = True
    for ar, en in GUIDE:
        guide.append([ar, en])
    guide["A1"].font = guide["B1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 80
    guide.column_dimensions["B"].width = 80

    _sheet(wb, "Courses", [k for k, _ in COLUMNS["Courses"]], [
        [c.id, c.name_ar, c.name, c.department, c.level, _yn(c.needs_lab)]
        for c in u.courses
    ])
    _sheet(wb, "Sections", [k for k, _ in COLUMNS["Sections"]], [
        [s.id, s.course_id, u.instructor_by_id[s.instructor_id].name, s.enrollment,
         ",".join(sorted(u.serving_cohorts.get(s.id, {s.cohort})))]
        for s in u.sections
    ])

    groups: dict[tuple, dict] = {}
    for st in u.students:
        key = (st.department, st.level, tuple(sorted(st.section_ids)), st.is_repeater)
        g = groups.get(key)
        if g is None:
            g = groups[key] = {"count": 0, "access": 0}
        g["count"] += 1
        g["access"] += st.needs_accessibility
    per_level: dict[tuple, int] = {}
    group_rows = []
    for (dept, level, sids, rep), g in groups.items():
        per_level[(dept, level)] = per_level.get((dept, level), 0) + 1
        group_rows.append([
            f"{dept}-L{level}-G{per_level[(dept, level)]}", dept, level, g["count"],
            ",".join(sids), _yn(rep), g["access"],
        ])
    _sheet(wb, "StudentGroups", [k for k, _ in COLUMNS["StudentGroups"]], group_rows)

    _sheet(wb, "Rooms", [k for k, _ in COLUMNS["Rooms"]], [
        [r.id, r.building_id, r.capacity, _yn(r.accessible), r.kind] for r in u.rooms
    ])
    bids = [b.id for b in u.buildings]
    _sheet(wb, "Buildings", [k for k, _ in COLUMNS["Buildings"]] + bids, [
        [b.id, b.name_ar, b.name] + [u.walk[(b.id, o)] for o in bids]
        for b in u.buildings
    ])
    _sheet(wb, "Instructors", [k for k, _ in COLUMNS["Instructors"]], [
        [i.name, i.name_en, format_unavailable(i.unavailable_slots)]
        for i in u.instructors
    ])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
