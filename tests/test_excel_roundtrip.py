"""The template must load back as the same campus, or the upload demo cannot
quote the same numbers as the built-in one."""

import io
from collections import Counter

from openpyxl import load_workbook as open_xlsx

from core.baseline import build_baseline
from core.metrics import compute
from data.excel_io import export_template, format_unavailable, load_workbook, parse_unavailable
from data.generator import generate


def _roundtrip():
    u = generate()
    v, report = load_workbook(export_template(u))
    return u, v, report


def _students(u):
    return Counter(
        (s.department, s.level, frozenset(s.section_ids), s.is_repeater, s.needs_accessibility)
        for s in u.students
    )


def test_template_loads_back_without_errors():
    _, v, report = _roundtrip()
    assert report["ok"], report["errors"]
    assert v is not None
    assert not report["warnings"]


def test_roundtrip_is_the_same_campus():
    u, v, _ = _roundtrip()
    assert len(v.sections) == len(u.sections)
    assert len(v.courses) == len(u.courses)
    assert len(v.rooms) == len(u.rooms)
    assert len(v.students) == len(u.students)
    assert {s.id: s.enrollment for s in v.sections} == {s.id: s.enrollment for s in u.sections}
    assert {s.id: (s.course_id, s.instructor_id) for s in v.sections} == {
        s.id: (s.course_id, s.instructor_id) for s in u.sections
    }
    assert {i.id: i.unavailable_slots for i in v.instructors} == {
        i.id: i.unavailable_slots for i in u.instructors
    }
    assert v.serving_cohorts == u.serving_cohorts
    assert v.walk == u.walk
    assert _students(v) == _students(u)


def test_roundtrip_gives_identical_baseline_metrics():
    u, v, _ = _roundtrip()
    a, b = compute(u, build_baseline(u)), compute(v, build_baseline(v))
    # the violation detail names student ids, which the workbook regenerates
    detail = "accessibility_violation_detail"
    assert {k: x for k, x in a.items() if k != detail} == {
        k: x for k, x in b.items() if k != detail
    }
    assert Counter(x["type"] for x in a[detail]) == Counter(x["type"] for x in b[detail])
    assert a["pct_students_with_2h_gap"] == 44.0


def test_same_file_same_campus():
    data = export_template(generate())
    v1, _ = load_workbook(data)
    v2, _ = load_workbook(data)
    assert [s.id for s in v1.students] == [s.id for s in v2.students]
    assert [s.section_ids for s in v1.students] == [s.section_ids for s in v2.students]


def test_unavailable_text_roundtrips():
    slots = parse_unavailable("Tue 14-16; Sun all; الخميس ٨-٩")
    assert len(slots) == 2 + 8 + 1
    assert parse_unavailable(format_unavailable(slots)) == slots


def _edit(data: bytes, fn) -> bytes:
    wb = open_xlsx(io.BytesIO(data))
    fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(wb, key):
    return next(ws for ws in wb.worksheets if key in ws.title)


def test_errors_carry_row_numbers_and_block_the_build():
    def break_it(wb):
        _sheet(wb, "Sections").cell(3, 2).value = "NOPE"        # unknown course
        _sheet(wb, "Rooms").cell(2, 3).value = 0                 # capacity 0
        _sheet(wb, "StudentGroups").cell(2, 5).value = "S01,S999"  # unknown section

    u, report = load_workbook(_edit(export_template(generate()), break_it))
    assert u is None and not report["ok"]
    where = {(e["sheet"], e["row"]) for e in report["errors"]}
    assert ("Sections", 3) in where
    assert ("Rooms", 2) in where
    assert ("StudentGroups", 2) in where
    assert all(e["ar"] for e in report["errors"])


def test_student_groups_are_optional():
    def drop(wb):
        wb.remove(_sheet(wb, "StudentGroups"))

    u, report = load_workbook(_edit(export_template(generate()), drop))
    assert report["ok"], report["errors"]
    assert report["notes"], "the derivation is explained in the report"
    # every level gets one group per section of its largest course, and a
    # group takes section g mod n of each course in its level
    for dept, level in {(c.department, c.level) for c in u.courses}:
        courses = [c.id for c in u.courses if (c.department, c.level) == (dept, level)]
        secs = {c: [s.id for s in u.sections if s.course_id == c] for c in courses}
        k = max(len(v) for v in secs.values())
        taken = {tuple(s.section_ids) for s in u.students
                 if (s.department, s.level) == (dept, level)}
        assert taken == {tuple(secs[c][g % len(secs[c])] for c in courses) for g in range(k)}


def test_cache_key_sees_more_than_section_ids():
    from api.state import campus_key

    u = generate()
    v, _ = load_workbook(export_template(u))
    assert campus_key(u) == campus_key(v), "a faithful round-trip shares the cache"
    v.sections[0].enrollment += 1
    assert campus_key(u) != campus_key(v), "same ids, different campus, different key"
